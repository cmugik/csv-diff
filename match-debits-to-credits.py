import os
import re
from tkinter import messagebox
from tkinter import filedialog
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import numpy as np


def load_csv_file():
    file_path = filedialog.askopenfilename(filetypes=[("CSV files", "*.csv")])
    if file_path:
        return pd.read_csv(file_path), file_path
    return None, None


def convert_and_sort(df, column_name):
    # remove dollar signs and commas, then convert to numeric
    df[column_name] = df[column_name].replace(
        '[$,]', '', regex=True).astype(float)
    sorted_df = df.sort_values(by=column_name, ascending=False)
    return sorted_df[column_name].values, sorted_df.index


def create_matchbooks(our_df, bank_df, our_columns, bank_columns, ours_is_credit):
    from datetime import datetime

    our_value_column, bank_value_column = (
        "Credits",
        "Debits",
    ) if ours_is_credit else ("Debits", "Credits")

    # --- Step 1: Clean numeric values ---
    our_df[our_value_column] = our_df[our_value_column].replace(
        "[$,]", "", regex=True
    ).astype(float)
    bank_df[bank_value_column] = bank_df[bank_value_column].replace(
        "[$,]", "", regex=True
    ).astype(float)

    # --- Step 2: Identify date columns ---
    possible_date_cols_our = [c for c in our_df.columns if "date" in c.lower()]
    possible_date_cols_bank = [
        c for c in bank_df.columns if "date" in c.lower()]
    our_date_col = possible_date_cols_our[0] if possible_date_cols_our else None
    bank_date_col = possible_date_cols_bank[0] if possible_date_cols_bank else None

    # --- Step 3: Parse dates explicitly as MM/DD/YYYY ---
    def parse_mmddyyyy(series):
        try:
            return pd.to_datetime(series, format="%m/%d/%Y", errors="raise")
        except Exception:
            return pd.to_datetime(series, format="%m/%d/%y", errors="coerce")

    if our_date_col and bank_date_col:
        our_df[our_date_col] = parse_mmddyyyy(our_df[our_date_col])
        bank_df[bank_date_col] = parse_mmddyyyy(bank_df[bank_date_col])

    # --- Step 4: Group by numeric value ---
    our_groups = {val: sub_df for val,
                  sub_df in our_df.groupby(our_value_column)}
    bank_groups = {val: sub_df for val,
                   sub_df in bank_df.groupby(bank_value_column)}

    results = []
    all_values = sorted(set(our_groups.keys()) | set(
        bank_groups.keys()), reverse=True)

    def fmt_date(dt):
        return dt.strftime("%m/%d/%Y") if pd.notnull(dt) else "XXX"

    # --- Step 5: Iterate through each value group ---
    for val in all_values:
        our_group = our_groups.get(val, pd.DataFrame())
        bank_group = bank_groups.get(val, pd.DataFrame())

        matched_bank_indices = set()
        matched_our_indices = set()

        # --- 5a: Match by value & exact date ---
        for our_idx, our_row in our_group.iterrows():
            our_date = our_row[our_date_col] if our_date_col else None
            for bank_idx, bank_row in bank_group.iterrows():
                bank_date = bank_row[bank_date_col] if bank_date_col else None
                if (
                    our_idx not in matched_our_indices
                    and bank_idx not in matched_bank_indices
                    and pd.notnull(our_date)
                    and pd.notnull(bank_date)
                    and our_date.date() == bank_date.date()
                ):
                    matched_our_indices.add(our_idx)
                    matched_bank_indices.add(bank_idx)

                    row = {
                        "Our_Value": val,
                        "Bank_Value": val,
                        "Match": "MATCH",
                    }
                    if our_date_col:
                        row[f"Our_{our_date_col}"] = fmt_date(
                            our_row[our_date_col])
                    if bank_date_col:
                        row[f"Bank_{bank_date_col}"] = fmt_date(
                            bank_row[bank_date_col])
                    for col in our_columns:
                        if col not in ["Credits", "Debits", our_date_col]:
                            row[f"Our_{col}"] = our_row[our_columns[col]]
                    for col in bank_columns:
                        if col not in ["Credits", "Debits", bank_date_col]:
                            row[f"Bank_{col}"] = bank_row[bank_columns[col]]
                    results.append(row)

        # --- 5b: Match remaining by value (ignore date) ---
        our_unmatched = our_group.loc[~our_group.index.isin(
            matched_our_indices)]
        bank_unmatched = bank_group.loc[~bank_group.index.isin(
            matched_bank_indices)]

        for our_idx, our_row in our_unmatched.iterrows():
            if bank_unmatched.empty:
                break
            bank_idx, bank_row = bank_unmatched.iloc[0].name, bank_unmatched.iloc[0]
            matched_our_indices.add(our_idx)
            matched_bank_indices.add(bank_idx)
            bank_unmatched = bank_unmatched.iloc[1:]

            row = {
                "Our_Value": val,
                "Bank_Value": val,
                "Match": "POTENTIAL_MATCH",
            }
            if our_date_col:
                row[f"Our_{our_date_col}"] = fmt_date(our_row[our_date_col])
            if bank_date_col:
                row[f"Bank_{bank_date_col}"] = fmt_date(
                    bank_row[bank_date_col])
            for col in our_columns:
                if col not in ["Credits", "Debits", our_date_col]:
                    row[f"Our_{col}"] = our_row[our_columns[col]]
            for col in bank_columns:
                if col not in ["Credits", "Debits", bank_date_col]:
                    row[f"Bank_{col}"] = bank_row[bank_columns[col]]
            results.append(row)

        # --- 5c: Remaining items = MISMATCH ---
        our_remaining = our_group.loc[~our_group.index.isin(
            matched_our_indices)]
        bank_remaining = bank_group.loc[~bank_group.index.isin(
            matched_bank_indices)]

        for _, our_row in our_remaining.iterrows():
            row = {
                "Our_Value": val,
                "Bank_Value": "XXX",
                "Match": "MISMATCH",
            }
            if our_date_col:
                row[f"Our_{our_date_col}"] = fmt_date(our_row[our_date_col])
            if bank_date_col:
                row[f"Bank_{bank_date_col}"] = "XXX"
            for col in our_columns:
                if col not in ["Credits", "Debits", our_date_col]:
                    row[f"Our_{col}"] = our_row[our_columns[col]]
            for col in bank_columns:
                if col not in ["Credits", "Debits", bank_date_col]:
                    row[f"Bank_{col}"] = "XXX"
            results.append(row)

        for _, bank_row in bank_remaining.iterrows():
            row = {
                "Our_Value": "XXX",
                "Bank_Value": val,
                "Match": "MISMATCH",
            }
            if our_date_col:
                row[f"Our_{our_date_col}"] = "XXX"
            if bank_date_col:
                row[f"Bank_{bank_date_col}"] = fmt_date(
                    bank_row[bank_date_col])
            for col in our_columns:
                if col not in ["Credits", "Debits", our_date_col]:
                    row[f"Our_{col}"] = "XXX"
            for col in bank_columns:
                if col not in ["Credits", "Debits", bank_date_col]:
                    row[f"Bank_{col}"] = bank_row[bank_columns[col]]
            results.append(row)

    return pd.DataFrame(results)


def save_to_excel_with_color(df, filename):
    wb = Workbook()
    ws = wb.active

    green_fill = PatternFill(start_color='00C851',
                             end_color='00C851', fill_type='solid')
    red_fill = PatternFill(start_color='FF4444',
                           end_color='FF4444', fill_type='solid')
    yellow_fill = PatternFill(start_color='FFEB3B',
                              end_color='FFEB3B', fill_type='solid')

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True)):
        ws.append(row)
        if r_idx == 0:
            continue
        match_status = row[2]
        if match_status == "MATCH":
            row_color = green_fill
        elif match_status == "POTENTIAL_MATCH":
            row_color = yellow_fill
        else:
            row_color = red_fill
        for cell in ws[r_idx + 1]:
            cell.fill = row_color

    ws.delete_cols(3)
    wb.save(filename)
    print("Saved " + filename)


class CSVMatcherApp:
    our_file_name = ""
    bank_file_name = ""
    file_month = ""

    def __init__(self, root):
        self.root = root
        self.root.title("CSV Matcher")

        self.our_csv = None
        self.bank_csv = None

        self.our_file_btn = tk.Button(
            root, text="Load OUR CSV", command=self.load_our_csv)
        self.our_file_btn.pack()

        self.bank_file_btn = tk.Button(
            root, text="Load BANK CSV", command=self.load_bank_csv)
        self.bank_file_btn.pack()

        self.match_btn = tk.Button(
            root, text="Match CSVs", command=self.match_csvs, state=tk.DISABLED)
        self.match_btn.pack()

        self.frame = tk.Frame(root)
        self.frame.pack(pady=20)

        self.our_col_label = tk.Label(self.frame, text="SAGE COLUMNS")
        self.our_col_label.grid(row=0, column=0, padx=5)

        self.bank_col_label = tk.Label(self.frame, text="BANK COLUMNS")
        self.bank_col_label.grid(row=0, column=2, padx=5)

        self.our_entries = []
        self.bank_entries = []

    def load_our_csv(self):
        self.our_csv, temp = load_csv_file()
        self.our_file_name = os.path.basename(temp)
        if self.our_csv is not None:
            month_pattern = r'\b(January|February|March|April|May|June|July|August|September|October|November|December)\b'
            match1 = re.search(
                month_pattern, self.our_file_name, re.IGNORECASE)
            match2 = re.search(
                month_pattern, self.bank_file_name, re.IGNORECASE)
            if match1 and match2 and match1.group(0) == match2.group(0):
                self.file_month = match1.group(0)
            self.display_columns(self.our_csv, 'our')
            self.check_files_loaded()

    def load_bank_csv(self):
        self.bank_csv, temp = load_csv_file()
        self.bank_file_name = os.path.basename(temp)
        if self.bank_csv is not None:
            month_pattern = r'\b(January|February|March|April|May|June|July|August|September|October|November|December)\b'
            match1 = re.search(
                month_pattern, self.our_file_name, re.IGNORECASE)
            match2 = re.search(
                month_pattern, self.bank_file_name, re.IGNORECASE)
            if match1 and match2 and match1.group(0) == match2.group(0):
                self.file_month = match1.group(0)
            self.display_columns(self.bank_csv, 'bank')
            self.check_files_loaded()

    def check_files_loaded(self):
        if self.our_csv is not None and self.bank_csv is not None:
            self.match_btn.config(state=tk.NORMAL)

    def match_csvs(self):
        our_columns = {entry.get(): entry.get()
                       for entry, _ in self.our_entries}
        bank_columns = {entry.get(): entry.get()
                        for entry, _ in self.bank_entries}

        matched_df = create_matchbooks(
            self.our_csv, self.bank_csv, our_columns, bank_columns, True)
        save_to_excel_with_color(
            matched_df, f"Matched_Output_OurCredits-TheirDebits_{self.file_month}.xlsx")

        matched_df = create_matchbooks(
            self.our_csv, self.bank_csv, our_columns, bank_columns, False)
        save_to_excel_with_color(
            matched_df, f"Matched_Output_OurDebits-TheirCredits_{self.file_month}.xlsx")
        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo(
            "Operation Complete", "The CSV matching operation has completed successfully!")
        root.destroy()

    def display_columns(self, df, col_type):
        columns = df.columns.tolist()

        # Filter out specific columns
        filtered_columns = [col for col in columns if col.lower(
        ) not in ['debits', 'credits', 'debit', 'credit']]

        if col_type == 'our':
            for entry, delete_btn in self.our_entries:
                entry.destroy()
                delete_btn.destroy()
            self.our_entries.clear()

            for idx, col in enumerate(filtered_columns):
                self.add_column_entry(col, idx, 'our')

        else:
            for entry, delete_btn in self.bank_entries:
                entry.destroy()
                delete_btn.destroy()
            self.bank_entries.clear()

            for idx, col in enumerate(filtered_columns):
                self.add_column_entry(col, idx, 'bank')

    def add_column_entry(self, col_name, idx, col_type):
        entry = tk.Entry(self.frame)
        entry.insert(0, col_name)
        entry.grid(row=idx + 1, column=0 if col_type ==
                   'our' else 2, padx=5, pady=5)

        delete_btn = tk.Button(self.frame, text="X", command=lambda: self.delete_entry(
            entry, delete_btn, col_type))
        delete_btn.grid(row=idx + 1, column=1 if col_type ==
                        'our' else 3, padx=5)

        if col_type == 'our':
            self.our_entries.append((entry, delete_btn))
        else:
            self.bank_entries.append((entry, delete_btn))

    def delete_entry(self, entry, delete_btn, col_type):
        if col_type == 'our':
            self.our_entries.remove((entry, delete_btn))
        else:
            self.bank_entries.remove((entry, delete_btn))

        entry.destroy()
        delete_btn.destroy()

        for idx, (e, btn) in enumerate(self.our_entries if col_type == 'our' else self.bank_entries):
            e.grid(row=idx + 1, column=0 if col_type ==
                   'our' else 2, padx=5, pady=5)
            btn.grid(row=idx + 1, column=1 if col_type == 'our' else 3, padx=5)


root = tk.Tk()
app = CSVMatcherApp(root)
root.mainloop()
