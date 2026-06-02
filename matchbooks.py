import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows


VALUE_COLUMNS = {"credits", "debits", "credit", "debit"}
MONTH_PATTERN = re.compile(
    r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\b",
    re.IGNORECASE,
)

DATE_FORMAT_OPTIONS = {
    "MM/DD/YYYY or MM/DD/YY": ("%m/%d/%Y", "%m/%d/%y"),
    "DD/MM/YYYY or DD/MM/YY": ("%d/%m/%Y", "%d/%m/%y"),
}

MATCH_MODES = {
    "Opposite sides (Credits/Debits)": (
        ("Credits", "Debits"),
        ("Debits", "Credits"),
    ),
    "Same sides (Debits/Debits, Credits/Credits)": (
        ("Debits", "Debits"),
        ("Credits", "Credits"),
    ),
}


def load_csv_file():
    file_path = filedialog.askopenfilename(filetypes=[("CSV files", "*.csv")])
    if file_path:
        return pd.read_csv(file_path), file_path
    return None, None


def find_date_column(columns):
    for col in columns:
        if "date" in col.lower():
            return col
    return None


def parse_dates(series, date_formats):
    for date_format in date_formats:
        parsed = pd.to_datetime(series, format=date_format, errors="coerce")
        if parsed.notna().any():
            return parsed
    return pd.to_datetime(series, errors="coerce")


def format_date(value):
    return value.strftime("%m/%d/%Y") if pd.notnull(value) else "XXX"


def add_sage_columns(row, sage_row, sage_columns, sage_date_col):
    for col in sage_columns:
        if col not in ["Credits", "Debits", sage_date_col]:
            row[f"Sage_{col}"] = sage_row[sage_columns[col]]


def add_bank_columns(row, bank_row, bank_columns, bank_date_col):
    for col in bank_columns:
        if col not in ["Credits", "Debits", bank_date_col]:
            row[f"Bank_{col}"] = bank_row[bank_columns[col]]


def add_empty_sage_columns(row, sage_columns, sage_date_col):
    for col in sage_columns:
        if col not in ["Credits", "Debits", sage_date_col]:
            row[f"Sage_{col}"] = "XXX"


def add_empty_bank_columns(row, bank_columns, bank_date_col):
    for col in bank_columns:
        if col not in ["Credits", "Debits", bank_date_col]:
            row[f"Bank_{col}"] = "XXX"


def create_matchbooks(
    sage_df,
    bank_df,
    sage_columns,
    bank_columns,
    sage_value_column,
    bank_value_column,
    date_formats=None,
):
    date_formats = date_formats or DATE_FORMAT_OPTIONS["MM/DD/YYYY or MM/DD/YY"]
    sage_df = sage_df.copy()
    bank_df = bank_df.copy()

    sage_df[sage_value_column] = (
        sage_df[sage_value_column].replace(
            "[$,]", "", regex=True).astype(float)
    )
    bank_df[bank_value_column] = (
        bank_df[bank_value_column].replace(
            "[$,]", "", regex=True).astype(float)
    )

    sage_df_filtered = sage_df[sage_df[sage_value_column] > 0].copy()
    bank_df_filtered = bank_df[bank_df[bank_value_column] > 0].copy()

    sage_date_col = find_date_column(sage_df_filtered.columns)
    bank_date_col = find_date_column(bank_df_filtered.columns)

    if sage_date_col and bank_date_col:
        sage_df_filtered[sage_date_col] = parse_dates(
            sage_df_filtered[sage_date_col], date_formats
        )
        bank_df_filtered[bank_date_col] = parse_dates(
            bank_df_filtered[bank_date_col], date_formats
        )

    sage_groups = {
        val: sub_df for val, sub_df in sage_df_filtered.groupby(sage_value_column)
    }
    bank_groups = {
        val: sub_df for val, sub_df in bank_df_filtered.groupby(bank_value_column)
    }

    results = []
    all_values = sorted(set(sage_groups.keys()) | set(
        bank_groups.keys()), reverse=True)

    for val in all_values:
        sage_group = sage_groups.get(val, pd.DataFrame())
        bank_group = bank_groups.get(val, pd.DataFrame())

        matched_bank_indices = set()
        matched_sage_indices = set()

        for sage_idx, sage_row in sage_group.iterrows():
            sage_date = sage_row[sage_date_col] if sage_date_col else None
            for bank_idx, bank_row in bank_group.iterrows():
                bank_date = bank_row[bank_date_col] if bank_date_col else None
                if (
                    sage_idx not in matched_sage_indices
                    and bank_idx not in matched_bank_indices
                    and pd.notnull(sage_date)
                    and pd.notnull(bank_date)
                    and sage_date.date() == bank_date.date()
                ):
                    matched_sage_indices.add(sage_idx)
                    matched_bank_indices.add(bank_idx)

                    row = {
                        "Sage_Value": val,
                        "Bank_Value": val,
                        "Match": "MATCH",
                    }
                    if sage_date_col:
                        row[f"Sage_{sage_date_col}"] = format_date(
                            sage_row[sage_date_col]
                        )
                    if bank_date_col:
                        row[f"Bank_{bank_date_col}"] = format_date(
                            bank_row[bank_date_col]
                        )
                    add_sage_columns(
                        row, sage_row, sage_columns, sage_date_col)
                    add_bank_columns(
                        row, bank_row, bank_columns, bank_date_col)
                    results.append(row)

        sage_unmatched = sage_group.loc[
            ~sage_group.index.isin(matched_sage_indices)
        ]
        bank_unmatched = bank_group.loc[~bank_group.index.isin(
            matched_bank_indices)]
        bank_unmatched_list = list(bank_unmatched.iterrows())

        for sage_idx, sage_row in sage_unmatched.iterrows():
            if not bank_unmatched_list:
                break

            sage_date = sage_row[sage_date_col] if sage_date_col else None
            best_match_idx = 0
            best_date_diff = float("inf")

            if sage_date_col and bank_date_col and pd.notnull(sage_date):
                for idx, (_, bank_row) in enumerate(bank_unmatched_list):
                    bank_date = bank_row[bank_date_col]
                    if pd.notnull(bank_date):
                        date_diff = abs((sage_date - bank_date).days)
                        if date_diff < best_date_diff:
                            best_date_diff = date_diff
                            best_match_idx = idx

            bank_idx, bank_row = bank_unmatched_list.pop(best_match_idx)
            matched_sage_indices.add(sage_idx)
            matched_bank_indices.add(bank_idx)

            row = {
                "Sage_Value": val,
                "Bank_Value": val,
                "Match": "POTENTIAL_MATCH",
            }
            if sage_date_col:
                row[f"Sage_{sage_date_col}"] = format_date(
                    sage_row[sage_date_col])
            if bank_date_col:
                row[f"Bank_{bank_date_col}"] = format_date(
                    bank_row[bank_date_col])
            add_sage_columns(row, sage_row, sage_columns, sage_date_col)
            add_bank_columns(row, bank_row, bank_columns, bank_date_col)
            results.append(row)

        sage_remaining = sage_group.loc[
            ~sage_group.index.isin(matched_sage_indices)
        ]
        bank_remaining = bank_group.loc[~bank_group.index.isin(
            matched_bank_indices)]

        for _, sage_row in sage_remaining.iterrows():
            row = {
                "Sage_Value": val,
                "Bank_Value": "XXX",
                "Match": "MISMATCH",
            }
            if sage_date_col:
                row[f"Sage_{sage_date_col}"] = format_date(
                    sage_row[sage_date_col])
            if bank_date_col:
                row[f"Bank_{bank_date_col}"] = "XXX"
            add_sage_columns(row, sage_row, sage_columns, sage_date_col)
            add_empty_bank_columns(row, bank_columns, bank_date_col)
            results.append(row)

        for _, bank_row in bank_remaining.iterrows():
            row = {
                "Sage_Value": "XXX",
                "Bank_Value": val,
                "Match": "MISMATCH",
            }
            if sage_date_col:
                row[f"Sage_{sage_date_col}"] = "XXX"
            if bank_date_col:
                row[f"Bank_{bank_date_col}"] = format_date(
                    bank_row[bank_date_col])
            add_empty_sage_columns(row, sage_columns, sage_date_col)
            add_bank_columns(row, bank_row, bank_columns, bank_date_col)
            results.append(row)

    return pd.DataFrame(results)


def save_to_excel_with_color(df, filename):
    wb = Workbook()
    ws = wb.active

    green_fill = PatternFill(
        start_color="00C851", end_color="00C851", fill_type="solid"
    )
    red_fill = PatternFill(start_color="FF4444",
                           end_color="FF4444", fill_type="solid")
    yellow_fill = PatternFill(
        start_color="FFEB3B", end_color="FFEB3B", fill_type="solid"
    )

    for row_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True)):
        ws.append(row)
        if row_idx == 0:
            continue

        match_status = row[2]
        if match_status == "MATCH":
            row_color = green_fill
        elif match_status == "POTENTIAL_MATCH":
            row_color = yellow_fill
        else:
            row_color = red_fill

        for cell in ws[row_idx + 1]:
            cell.fill = row_color

    ws.delete_cols(3)
    wb.save(filename)
    print("Saved " + filename)


class CSVMatcherApp:
    def __init__(self, root):
        self.root = root
        self.root.title("CSV Matchbooks")

        self.sage_csv = None
        self.bank_csv = None
        self.sage_file_name = ""
        self.bank_file_name = ""
        self.file_month = ""

        self.sage_entries = []
        self.bank_entries = []

        self.top_frame = tk.Frame(root)
        self.top_frame.pack(padx=12, pady=12, fill=tk.X)

        self.sage_file_btn = tk.Button(
            self.top_frame, text="Load SAGE CSV", command=self.load_sage_csv
        )
        self.sage_file_btn.grid(row=0, column=0, padx=5, pady=5, sticky="ew")

        self.bank_file_btn = tk.Button(
            self.top_frame, text="Load BANK CSV", command=self.load_bank_csv
        )
        self.bank_file_btn.grid(row=0, column=1, padx=5, pady=5, sticky="ew")

        tk.Label(self.top_frame, text="Match mode").grid(
            row=1, column=0, padx=5, pady=5, sticky="w"
        )
        self.match_mode = tk.StringVar(value=next(iter(MATCH_MODES)))
        self.match_mode_menu = tk.OptionMenu(
            self.top_frame, self.match_mode, *MATCH_MODES.keys()
        )
        self.match_mode_menu.grid(row=1, column=1, padx=5, pady=5, sticky="ew")

        tk.Label(self.top_frame, text="Date format").grid(
            row=2, column=0, padx=5, pady=5, sticky="w"
        )
        self.date_format = tk.StringVar(value=next(iter(DATE_FORMAT_OPTIONS)))
        self.date_format_menu = tk.OptionMenu(
            self.top_frame, self.date_format, *DATE_FORMAT_OPTIONS.keys()
        )
        self.date_format_menu.grid(
            row=2, column=1, padx=5, pady=5, sticky="ew")

        self.match_btn = tk.Button(
            self.top_frame,
            text="Match CSVs",
            command=self.match_csvs,
            state=tk.DISABLED,
        )
        self.match_btn.grid(row=3, column=0, columnspan=2,
                            padx=5, pady=5, sticky="ew")

        self.status_var = tk.StringVar(value="Load both CSVs to begin.")
        self.status_label = tk.Label(
            self.top_frame, textvariable=self.status_var, fg="gray20"
        )
        self.status_label.grid(
            row=4, column=0, columnspan=2, padx=5, pady=(2, 0), sticky="w"
        )

        self.columns_frame = tk.Frame(root)
        self.columns_frame.pack(padx=12, pady=(
            0, 12), fill=tk.BOTH, expand=True)

        self.sage_frame = tk.LabelFrame(
            self.columns_frame, text="SAGE COLUMNS")
        self.sage_frame.grid(row=0, column=0, padx=6, sticky="nsew")

        self.bank_frame = tk.LabelFrame(
            self.columns_frame, text="BANK COLUMNS")
        self.bank_frame.grid(row=0, column=1, padx=6, sticky="nsew")

        self.columns_frame.columnconfigure(0, weight=1)
        self.columns_frame.columnconfigure(1, weight=1)
        self.top_frame.columnconfigure(1, weight=1)

    def load_sage_csv(self):
        self.sage_csv, file_path = load_csv_file()
        if self.sage_csv is not None:
            self.sage_file_name = os.path.basename(file_path)
            self.set_file_month()
            self.display_columns(self.sage_csv, "sage")
            self.check_files_loaded()

    def load_bank_csv(self):
        self.bank_csv, file_path = load_csv_file()
        if self.bank_csv is not None:
            self.bank_file_name = os.path.basename(file_path)
            self.set_file_month()
            self.display_columns(self.bank_csv, "bank")
            self.check_files_loaded()

    def set_file_month(self):
        sage_match = MONTH_PATTERN.search(self.sage_file_name)
        bank_match = MONTH_PATTERN.search(self.bank_file_name)
        if sage_match and bank_match and sage_match.group(0) == bank_match.group(0):
            self.file_month = sage_match.group(0)

    def check_files_loaded(self):
        if self.sage_csv is not None and self.bank_csv is not None:
            self.match_btn.config(state=tk.NORMAL)
            self.status_var.set("Ready to match.")

    def match_csvs(self):
        sage_columns = self.collect_columns(self.sage_entries)
        bank_columns = self.collect_columns(self.bank_entries)
        date_formats = DATE_FORMAT_OPTIONS[self.date_format.get()]
        match_pairs = MATCH_MODES[self.match_mode.get()]
        output_files = []

        self.set_processing_state(True)
        try:
            for sage_value_column, bank_value_column in match_pairs:
                matched_df = create_matchbooks(
                    self.sage_csv,
                    self.bank_csv,
                    sage_columns,
                    bank_columns,
                    sage_value_column,
                    bank_value_column,
                    date_formats,
                )
                filename = (
                    f"Matched_Output_Sage{sage_value_column}-"
                    f"Bank{bank_value_column}_{self.file_month}.xlsx"
                )
                save_to_excel_with_color(matched_df, filename)
                output_files.append(filename)
        except Exception as exc:
            self.status_var.set("Matching failed.")
            messagebox.showerror("Matching Failed", str(exc))
        else:
            self.status_var.set("Matching complete.")
            messagebox.showinfo(
                "Operation Complete",
                "The CSV matching operation has completed successfully!\n\n"
                + "\n".join(output_files),
            )
        finally:
            self.set_processing_state(False)

    def set_processing_state(self, is_processing):
        state = tk.DISABLED if is_processing else tk.NORMAL
        self.sage_file_btn.config(state=state)
        self.bank_file_btn.config(state=state)
        self.match_btn.config(state=state)
        self.match_mode_menu.config(state=state)
        self.date_format_menu.config(state=state)
        if is_processing:
            self.status_var.set("Processing... please wait.")
            self.root.update_idletasks()

    def display_columns(self, df, col_type):
        entries = self.sage_entries if col_type == "sage" else self.bank_entries
        frame = self.sage_frame if col_type == "sage" else self.bank_frame

        for child in frame.winfo_children():
            child.destroy()
        entries.clear()

        columns = df.columns.tolist()
        date_col = find_date_column(columns)
        visible_columns = [
            col
            for col in columns
            if col.lower() not in VALUE_COLUMNS and col != date_col
        ]

        row_idx = 0
        tk.Label(frame, text="Date column", font=("TkDefaultFont", 9, "bold")).grid(
            row=row_idx, column=0, padx=5, pady=(8, 2), sticky="w"
        )
        row_idx += 1

        if date_col:
            self.add_column_entry(date_col, row_idx, col_type, is_date=True)
        else:
            tk.Label(frame, text="No date column detected", fg="gray40").grid(
                row=row_idx, column=0, columnspan=2, padx=5, pady=5, sticky="w"
            )
        row_idx += 1

        tk.Label(frame, text="Other output columns").grid(
            row=row_idx, column=0, padx=5, pady=(10, 2), sticky="w"
        )
        row_idx += 1

        for col in visible_columns:
            self.add_column_entry(col, row_idx, col_type, is_date=False)
            row_idx += 1

    def add_column_entry(self, col_name, row_idx, col_type, is_date=False):
        frame = self.sage_frame if col_type == "sage" else self.bank_frame
        entries = self.sage_entries if col_type == "sage" else self.bank_entries

        entry = tk.Entry(frame)
        entry.insert(0, col_name)
        entry.grid(row=row_idx, column=0, padx=5, pady=4, sticky="ew")
        frame.columnconfigure(0, weight=1)

        if is_date:
            entry.config(state="readonly", readonlybackground="#fff3bf")
            delete_btn = None
        else:
            delete_btn = tk.Button(
                frame,
                text="X",
                command=lambda: self.delete_entry(entry, delete_btn, entries),
            )
            delete_btn.grid(row=row_idx, column=1, padx=5, pady=4)

        entries.append(
            {"entry": entry, "delete_btn": delete_btn, "is_date": is_date})

    def delete_entry(self, entry, delete_btn, entries):
        matching_entry = next(
            item for item in entries if item["entry"] == entry and item["delete_btn"] == delete_btn
        )
        if matching_entry["is_date"]:
            return

        entries.remove(matching_entry)
        entry.destroy()
        delete_btn.destroy()

    def collect_columns(self, entries):
        return {item["entry"].get(): item["entry"].get() for item in entries}


def main():
    root = tk.Tk()
    CSVMatcherApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
